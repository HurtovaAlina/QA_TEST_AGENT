# Adds generated test cases to Excel file (new test cases → Excel)
import openpyxl

# Generates ID for new test case
def get_next_test_case_id(sheet):
    max_id = 0

    for row in sheet.iter_rows(values_only=True):
        test_id = row[0]

        if isinstance(test_id, str) and test_id.startswith("AI-"):
            try:
                number = int(test_id.replace("AI-", ""))

                if number > max_id:
                    max_id = number

            except ValueError: #
                continue

    return f"AI-{max_id + 1:03d}"

# Adds test case to Excel
def add_test_cases_to_excel(file_path: str, sheet_name: str, test_cases):
    """
    Adds generated test cases to Excel file.
    :param file_path: Path to Excel file
    :param sheet_name: Name of the sheet to add test cases
    :param test_cases: List of test cases to add
    """

    # open file with test cases
    workbook = openpyxl.load_workbook(file_path)

    if sheet_name not in workbook.sheetnames:
        sheet = workbook.create_sheet(sheet_name)

        sheet["A1"] = "ID"
        sheet["B1"] = "Title"
        sheet["C1"] = "Test Step"
        sheet["D1"] = "Step Action"
        sheet["E1"] = "Step Expected"

    else:
        sheet = workbook[sheet_name]

    # add test cases to sheet
    for test_case in test_cases:

        test_case_id = get_next_test_case_id(sheet)

        for index, step in enumerate(test_case["steps"]):

            row = sheet.max_row + 1

            if index == 0:
                sheet.cell(
                    row=row,
                    column=1,
                    value=test_case_id
                )

                sheet.cell(
                    row=row,
                    column=2,
                    value=test_case["test_case_name"]
                )

            sheet.cell(
                row=row,
                column=3,
                value=step["test_step"]
            )

            sheet.cell(
                row=row,
                column=4,
                value=""
            )

            sheet.cell(
                row=row,
                column=5,
                value=step["expected_result"]
            )

        # Empty row between test cases
        empty_row = sheet.max_row + 1
        sheet.cell(row=empty_row, column=1, value="")

    # Saves changes to the file
    workbook.save(file_path)