
from docx import Document as DocxDocument # to read docx files
from langchain_core.documents import Document as LangChainDocument # to save document as object
# with additional data - page content(text) and metadata(feature name)

def load_document(file_path:str)-> list[str]:

    """
    Load a docx file and return a list of paragraphs.
    :param file_path: path to  .docx file
    :return: paragraphs - list of paragraphs
    """

    doc = DocxDocument(file_path)

    paragraphs = []

    for paragraph in doc.paragraphs:
        paragraphs.append(paragraph.text.strip())

    return paragraphs


def split_into_features(paragraphs: list[str])-> list[LangChainDocument]:

    """
    Split a list of paragraphs into structured features. Features are separated by 2 empty paragraphs.
    :param paragraphs: list of paragraphs
    :return: features - list of features. Each feature is an object  of LangChainDocument.
    This document contains text for current feature and metadata with feature name.
    """
    features = []

    current_feature = None
    current_text = [] # text for current feature

    empty_paragraphs = 0

    for paragraph in paragraphs:

        # Empty paragraph
        if paragraph == "":
            empty_paragraphs += 1
            continue

        # If found text after 2 empty paragraphs - it means a new feature found
        if empty_paragraphs >= 2:

            if current_feature is not None:
                features.append(
                    LangChainDocument(
                        page_content="\n".join(current_text),
                        metadata={
                            "feature": current_feature
                        }
                    )
                )

            current_feature = paragraph
            current_text = []

        else:
            # First text in the document
            if current_feature is None:
                current_feature = paragraph
            else: # or text for current feature
                current_text.append(paragraph)

        empty_paragraphs = 0 # counter of empty paragraphs is reset

    # Added last feature
    if current_feature is not None:
        features.append(
            LangChainDocument(
                page_content="\n".join(current_text),
                metadata={
                    "feature": current_feature
                }
            )
        )

    return features